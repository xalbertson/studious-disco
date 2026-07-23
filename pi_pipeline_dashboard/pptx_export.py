"""Export exhibits as unformatted PPTX slides with a native, editable Excel-backed chart.

Deliberately left unstyled (default python-pptx template, default chart colors) -
these are working exhibits meant to be pasted into / restyled within a firm deck,
not a finished presentation.
"""

import datetime as dt
import io

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

_BLANK_LAYOUT_INDEX = 6  # "Blank" layout in python-pptx's default template


def build_filter_footnote(filters: dict, n: int) -> str:
    """filters: label -> value (str, list, or falsy/None to omit)."""
    parts = []
    for label, value in filters.items():
        if not value:
            continue
        if isinstance(value, (list, tuple, set)):
            parts.append(f"{label} = {', '.join(str(v) for v in value)}")
        else:
            parts.append(f"{label} = {value}")
    filter_str = "; ".join(parts) if parts else "None"
    today = dt.date.today().isoformat()
    return f"Filters: {filter_str}  |  n = {n} firms  |  Generated {today}"


def build_bar_chart_pptx(
    title: str,
    categories: list,
    series: dict,
    footnote: str,
    stacked: bool = False,
) -> bytes:
    """series: dict of series name -> list of values, aligned to `categories`."""
    chart_type = XL_CHART_TYPE.COLUMN_STACKED if stacked else XL_CHART_TYPE.COLUMN_CLUSTERED

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[_BLANK_LAYOUT_INDEX])

    title_box = slide.shapes.add_textbox(
        Inches(0.5), Inches(0.3), prs.slide_width - Inches(1), Inches(0.6)
    )
    title_box.text_frame.text = title
    title_box.text_frame.paragraphs[0].font.size = Pt(24)
    title_box.text_frame.paragraphs[0].font.bold = True

    chart_data = CategoryChartData()
    chart_data.categories = categories
    for name, values in series.items():
        chart_data.add_series(name, values)

    chart_top = Inches(1.1)
    chart_height = prs.slide_height - chart_top - Inches(1.0)
    slide.shapes.add_chart(
        chart_type,
        Inches(0.5),
        chart_top,
        prs.slide_width - Inches(1),
        chart_height,
        chart_data,
    )

    footnote_box = slide.shapes.add_textbox(
        Inches(0.5),
        prs.slide_height - Inches(0.6),
        prs.slide_width - Inches(1),
        Inches(0.4),
    )
    footnote_tf = footnote_box.text_frame
    footnote_tf.word_wrap = True
    footnote_tf.text = footnote
    footnote_tf.paragraphs[0].font.size = Pt(10)
    footnote_tf.paragraphs[0].font.italic = True

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
